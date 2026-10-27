from __future__ import annotations

import datetime
import importlib.util
import logging
import shutil
import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

RETRY_ATTEMPTS = 3
RETRY_DELAY_SECONDS = 3.0
EXCEL_ENGINE = "xlsxwriter" if importlib.util.find_spec("xlsxwriter") else "openpyxl"


def stringify_date_time_columns(df: pd.DataFrame) -> pd.DataFrame:
    for col in df.columns:
        sample = df[col].dropna()
        if not sample.empty and isinstance(sample.iloc[0], (datetime.date, datetime.time)):
            df[col] = df[col].astype(str).replace(["NaT", "None", "nan"], "")
    return df


def to_excel_with_retry(df: pd.DataFrame, path: Path, attempts: int = RETRY_ATTEMPTS,
                        delay_seconds: float = RETRY_DELAY_SECONDS) -> bool:
    for attempt in range(1, attempts + 1):
        try:
            df.to_excel(path, index=False, engine=EXCEL_ENGINE)
            return True
        except PermissionError:
            if attempt < attempts:
                logger.warning("File %s đang bị khóa. Thử lại sau %ss...", path.name, delay_seconds)
                time.sleep(delay_seconds)
            else:
                logger.warning("Không thể ghi file %s sau %s lần thử (file đang mở).", path, attempts)
        except Exception as exc:
            logger.warning("Lỗi ghi file %s: %s", path, exc)
            return False
    return False


def load_workbook_with_retry(path: Path, attempts: int = RETRY_ATTEMPTS, delay_seconds: float = RETRY_DELAY_SECONDS):
    for attempt in range(1, attempts + 1):
        try:
            return load_workbook(path)
        except PermissionError:
            if attempt < attempts:
                time.sleep(delay_seconds)
            else:
                logger.warning("Không thể mở file FEFO %s (đang bị khóa).", path)
        except Exception as exc:
            logger.warning("Lỗi mở file FEFO %s: %s", path, exc)
            return None
    return None


def save_workbook_with_retry(wb, path: Path, attempts: int = RETRY_ATTEMPTS,
                             delay_seconds: float = RETRY_DELAY_SECONDS) -> bool:
    for attempt in range(1, attempts + 1):
        try:
            wb.save(path)
            return True
        except PermissionError:
            if attempt < attempts:
                time.sleep(delay_seconds)
            else:
                logger.warning("Không thể lưu file FEFO %s (đang bị khóa).", path)
        except Exception as exc:
            logger.warning("Lỗi lưu file FEFO %s: %s", path, exc)
            return False
    return False


def copy_with_retry(source: Path, target: Path, attempts: int = RETRY_ATTEMPTS,
                    delay_seconds: float = RETRY_DELAY_SECONDS) -> bool:
    for attempt in range(1, attempts + 1):
        try:
            shutil.copyfile(source, target)
            return True
        except PermissionError:
            if attempt < attempts:
                time.sleep(delay_seconds)
            else:
                logger.warning("Không thể ghi file %s sau %s lần thử (file đang mở).", target, attempts)
        except Exception as exc:
            logger.warning("Lỗi ghi file %s: %s", target, exc)
            return False
    return False
