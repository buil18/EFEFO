from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl.cell.cell import MergedCell

from efefo.infra.files import load_workbook_with_retry, save_workbook_with_retry

logger = logging.getLogger(__name__)


def unmerge_overlapping_ranges(ws, min_row: int, max_row: int) -> None:
    ranges = [str(r) for r in list(ws.merged_cells.ranges) if r.min_row <= max_row and r.max_row >= min_row]
    for rng in ranges:
        ws.unmerge_cells(rng)


def export_fefo_workbook(summary: pd.DataFrame | None, fefo_file: Path, sheet_name: str, start_row: int) -> None:
    if not fefo_file.exists():
        logger.warning("Không tìm thấy file FEFO tại %s — bỏ qua.", fefo_file)
        return
    if summary is None:
        return

    wb = load_workbook_with_retry(fefo_file)
    if wb is None or sheet_name not in wb.sheetnames:
        return

    ws = wb[sheet_name]
    old_max_row = max(ws.max_row, start_row)
    new_max_row = start_row + len(summary) - 1 if len(summary) else start_row
    unmerge_overlapping_ranges(ws, start_row, max(old_max_row, new_max_row))

    for row in ws.iter_rows(min_row=start_row, max_row=old_max_row):
        for cell in row:
            if not isinstance(cell, MergedCell):
                cell.value = None

    for i, record in enumerate(summary.itertuples(index=False), start=start_row):
        for j, value in enumerate(record, start=1):
            cell = ws.cell(row=i, column=j)
            if not isinstance(cell, MergedCell):
                cell.value = value

    if save_workbook_with_retry(wb, fefo_file):
        logger.info("Đã cập nhật %s dòng vào sheet '%s' của file FEFO.", len(summary), sheet_name)
