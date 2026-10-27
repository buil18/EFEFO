from __future__ import annotations

import gzip
import importlib.util
import json
import os
import sys
from contextlib import contextmanager
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from efefo.config import load_site
from efefo.context import build_context
from efefo.runner import Services, run
from tests.support.fakes import FakeHttp

FEFO_URL = "https://script.google.com/macros/s/FAKE-FEFO/exec"
SECRET = "test-secret"


@contextmanager
def environment(**values):
    saved = {k: os.environ.get(k) for k in values}
    os.environ.update({k: v for k, v in values.items() if v is not None})
    for k, v in values.items():
        if v is None:
            os.environ.pop(k, None)
    try:
        yield
    finally:
        for k, v in saved.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v


def base_env(output_dir: Path, session_dir: Path) -> dict:
    return {
        "SAP_OUTPUT_DIR": str(output_dir),
        "SAP_SESSION_DIR": str(session_dir),
        "SAP_INPUT_DIR": None,
        "GOOGLE_APPS_SCRIPT_URL": FEFO_URL,
        "GOOGLE_APPS_SCRIPT_SECRET": SECRET,
        "GOOGLE_CAPACITY_SCRIPT_URL": None,
        "SAP_TEST_MODE": None,
        "SAP_MA_BO_SUNG": None,
        "SAP_PRD_DATE_FORMAT": None,
        "SAP_URL": None,
        "SAP_UNATTENDED": "1",
    }


def load_legacy(path: str):
    spec = importlib.util.spec_from_file_location("legacy_sap_full", path)
    module = importlib.util.module_from_spec(spec)
    sys.modules["legacy_sap_full"] = module
    spec.loader.exec_module(module)
    return module


def run_legacy(legacy, output_dir: Path, session_dir: Path, records: list[dict], http: FakeHttp) -> None:
    legacy.fetch_all_records = lambda cfg: records
    legacy._requests_session_with_retry = http.factory
    with environment(**base_env(output_dir, session_dir)):
        legacy.main()


def run_new(output_dir: Path, session_dir: Path, records: list[dict], http: FakeHttp, site=None):
    site = site or load_site("hm")
    with environment(**base_env(output_dir, session_dir)):
        ctx = build_context(site)
        return run(ctx, Services(fetch_records=lambda c: records, http_session=http.factory))


def read_excel_all(path: Path) -> dict[str, pd.DataFrame]:
    return pd.read_excel(path, sheet_name=None, dtype=object)


def workbook_values(path: Path) -> dict[str, list[tuple]]:
    wb = load_workbook(path)
    return {ws.title: [tuple(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def gz_json(path: Path):
    with gzip.open(path, "rt", encoding="utf-8") as handle:
        return json.load(handle)


def snapshot_outputs(root: Path) -> dict:
    result = {}
    for path in sorted(root.rglob("*")):
        if path.is_dir():
            continue
        rel = path.relative_to(root).as_posix()
        if rel.startswith("logs/") or rel.endswith(".lock"):
            continue
        if path.suffix == ".xlsx":
            result[rel] = workbook_values(path) if path.name == "FEFO.xlsx" else read_excel_all(path)
        elif path.name.endswith(".json.gz"):
            result[rel] = gz_json(path)
        elif path.suffix == ".json":
            result[rel] = json.loads(path.read_text(encoding="utf-8"))
        elif path.suffix == ".jsonl":
            result[rel] = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
        else:
            result[rel] = path.read_text(encoding="utf-8")
    return result


def assert_snapshots_equal(old: dict, new: dict) -> None:
    assert sorted(old) == sorted(new), f"Khác danh sách file: {sorted(set(old) ^ set(new))}"
    for key in old:
        a, b = old[key], new[key]
        if isinstance(a, dict) and a and all(isinstance(v, pd.DataFrame) for v in a.values()):
            assert sorted(a) == sorted(b), key
            for sheet in a:
                pd.testing.assert_frame_equal(a[sheet], b[sheet], check_dtype=True, obj=f"{key}:{sheet}")
        else:
            assert a == b, f"Khác nội dung: {key}"
