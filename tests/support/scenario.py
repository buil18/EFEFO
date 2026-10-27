from __future__ import annotations

import csv
import datetime
import io
import random
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import Workbook

MATERIALS = [
    ("10128309", "HC SILVER 250", "HEINEKEN Silver Can Case 24x25", "CS", 169, 365),
    ("10079475", "TP", "TIGER RB Crate 24x33cl VN", "CR", 70, 365),
    ("10079464", "TP CY", "TIGER Crystal RB Crate 24x33cl", "CR", 70, 365),
    ("10079514", "HC SLEEK", "HEINEKEN Can Case 24x33cl SLK", "CS", 130, 365),
    ("10136266", "TC SMOOTH 330", "TIGER Smooth Can 330", "CS", 120, 270),
    ("10090001", "LARUE", "LARUE Can Case 24x33", "TRAY", 150, 180),
    ("10090002", "KEG HNK", "HEINEKEN Keg 30L", "KEG", 12, 120),
    ("10090003", "BIVINA", "BIVINA Bottle Crate", "CR", 60, 0),
]
UNKNOWN_MATERIAL = "10999999"
MANUAL_ONLY = ["10777001", "10777002"]
MANUAL_ONLY_MASTER = [
    ("10777001", "NEW SKU A", "New SKU A Can", "CS", 100, 300),
    ("10777002", "NEW SKU B", "New SKU B Crate", "CR", 50, 300),
]
ZONES = ["D", "E", "F", "G", "BL", "K", "H"]


def odata_date(day: datetime.date | None) -> str | None:
    if day is None:
        return None
    epoch = datetime.datetime(1970, 1, 1)
    return f"/Date({int((datetime.datetime.combine(day, datetime.time()) - epoch).total_seconds() * 1000)})/"


def raw_bin(rng: random.Random, base: str) -> str:
    style = rng.random()
    if style < 0.55:
        return f"HM{base}"
    if style < 0.65:
        return f"HMPKG{base}"
    if style < 0.72:
        return f"PKG{base}"
    if style < 0.80:
        return f"hm {base.lower()}"
    if style < 0.86:
        return f"HM{base}  "
    return base


@dataclass
class Scenario:
    today: datetime.date
    records: list[dict]
    stock_count_csv: str
    manual_materials: list[str]
    locations: list[str] = field(default_factory=list)


def location_pool(rng: random.Random, size: int) -> list[str]:
    pool = []
    for zone in ZONES:
        for n in range(1, size + 1):
            pool.append(f"{zone}{n:02d}")
    pool += ["LOOSE-01", "FDD1", "BL7.3", "E29", "F31"]
    rng.shuffle(pool)
    return pool


def build_records(rng: random.Random, today: datetime.date, locations: list[str], n_records: int) -> list[dict]:
    records = []
    for i in range(n_records):
        material = rng.choice([m[0] for m in MATERIALS] + [UNKNOWN_MATERIAL])
        base = rng.choice(locations)
        age = rng.choice([0, 1, 2, 3, 5, 6, 7, 8, 10, 14, 30, 60, 120, 200, 300, 400])
        prod = today - datetime.timedelta(days=age)
        entry = prod + datetime.timedelta(days=rng.choice([0, 0, 1, 2]))
        prod_value = odata_date(prod) if rng.random() > 0.05 else None
        entry_value = odata_date(entry) if rng.random() > 0.03 else None
        qty = rng.choice([0, 1, 50, 120, 169, 250, 338, 500, 1000, 12.5, 0.333, 33.125, 7.1, 101.7])
        records.append({
            "__metadata": {"id": f"https://sap/ZIM_BCTKTTT_MAIN('{i % (n_records - 3) if n_records > 10 else i}')"},
            "Material": material,
            "Batch": f"L{rng.randint(1000000, 9999999)}F",
            "EWMWarehouse": "HM1",
            "EWMStorageType": rng.choice(["Y001", "Y002", "Y910"]),
            "EWMStorageBin": raw_bin(rng, base) if rng.random() > 0.01 else "",
            "NgaySanXuat": prod_value,
            "NgayNhapKho": entry_value,
            "ThoiGianSanXuat": f"PT{rng.randint(0, 23):02d}H{rng.randint(0, 59):02d}M{rng.randint(0, 59):02d}S",
            "EWMStockQuantityInBaseUnit": f"{qty:.3f}",
            "Unrestricted": f"{qty:.3f}",
            "Blocked": "0.000",
            "HandlingUnitNumber": f"{rng.randint(10**8, 10**9)}",
            "ProductType": "FERT",
            "BaseUnit": "CS",
        })
    return records


def build_stock_count(rng: random.Random, today: datetime.date, records: list[dict], locations: list[str],
                      n_rows: int) -> str:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(["KIEM KHO THUC TE", "", "", "", "", "", ""])
    writer.writerow(["Location", "Item", "SKU", "PRD date", "Qty (PL)", "Q'TY (Carton/Crate)", "STATUS"])
    sap_rows = [r for r in records if r["EWMStorageBin"] and r["NgaySanXuat"]]
    for _ in range(n_rows):
        kind = rng.random()
        if kind < 0.45 and sap_rows:
            src = rng.choice(sap_rows)
            ms = int(src["NgaySanXuat"][6:-2])
            day = (datetime.datetime(1970, 1, 1) + datetime.timedelta(milliseconds=ms)).date()
            loc = src["EWMStorageBin"].replace("HMPKG", "").replace("PKG", "").strip()
            material = src["Material"]
        elif kind < 0.75:
            day = today - datetime.timedelta(days=rng.choice([1, 4, 9, 40, 100]))
            loc = rng.choice(locations)
            material = rng.choice(MANUAL_ONLY)
        else:
            day = today - datetime.timedelta(days=rng.choice([2, 20, 90]))
            loc = rng.choice(locations)
            material = rng.choice([m[0] for m in MATERIALS])
        style = rng.random()
        if style < 0.6:
            prd = day.strftime("%d/%m/%Y")
        elif style < 0.75:
            prd = day.strftime("%Y-%m-%d")
        elif style < 0.85:
            prd = day.strftime("%d/%m/%y")
        elif style < 0.95:
            prd = day.strftime("%d-%m-%Y")
        else:
            prd = ""
        loc_text = rng.choice([loc, f" {loc.lower()} ", f"HM{loc}", f"{loc}​"]) if rng.random() < 0.3 else loc
        qty_pl = rng.choice(["0", "1", "2", "0.5", "3", ""])
        carton = rng.choice(["", "", "6.760", "120", "1,250", "0", "85"])
        status = rng.choice(["", "", "", "", "Blocked", "block", "OK"])
        writer.writerow([loc_text, material + (".0" if rng.random() < 0.1 else ""), "SKU " + material[-3:], prd,
                         qty_pl, carton, status])
    return out.getvalue()


def build_scenario(seed: int, today: datetime.date, n_records: int = 400, n_stock_rows: int = 120,
                   pool_size: int = 12) -> Scenario:
    rng = random.Random(seed)
    locations = location_pool(rng, pool_size)
    records = build_records(rng, today, locations, n_records)
    csv_text = build_stock_count(rng, today, records, locations, n_stock_rows)
    return Scenario(today=today, records=records, stock_count_csv=csv_text, manual_materials=MANUAL_ONLY,
                    locations=locations)


def mutate_records(records: list[dict], seed: int, drop_ratio: float = 0.05) -> list[dict]:
    rng = random.Random(seed)
    result = []
    for record in records:
        if rng.random() < drop_ratio:
            continue
        clone = dict(record)
        if rng.random() < 0.25:
            qty = float(clone["EWMStockQuantityInBaseUnit"])
            clone["EWMStockQuantityInBaseUnit"] = f"{max(qty - rng.choice([50, 150, 300, 0.1, 101.3]), 0):.3f}"
        result.append(clone)
    return result


def write_shelf_life(path: Path, include_manual: bool = True) -> None:
    rows = MATERIALS + (MANUAL_ONLY_MASTER if include_manual else [])
    pd.DataFrame(
        [{"Iteam": int(m[0]), "Short Name": m[1], "ID10126716": int(m[0]), "Name": m[2], "Unit": m[3],
          "Qty (Carton/Crate)/Pallet": m[4], "Local Shelf Life": m[5]} for m in rows]
    ).to_excel(path, index=False)


def write_capacity(path: Path, locations: list[str]) -> None:
    cap = pd.DataFrame([
        {"Group location": loc[0], "Location": loc, "Capacity Can": 52, "Capacity Bot": 54 if i % 7 else 0,
         "Capacity Keg": 108, "Note": 1}
        for i, loc in enumerate(sorted(set(locations)) + ["Z99", "DDM1"])
    ])
    unit = pd.DataFrame([{"Material": int(m[0]), "Name": m[2], "Unit": m[3]}
                         for m in MATERIALS + MANUAL_ONLY_MASTER])
    with pd.ExcelWriter(path) as writer:
        cap.to_excel(writer, sheet_name="Capacity", index=False)
        unit.to_excel(writer, sheet_name="Unit", index=False)


def write_fefo_template(path: Path, sheet: str = "Update Actual Count", stale_rows: int = 30) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws["A1"] = "FEFO"
    ws.merge_cells("A1:F1")
    ws["A3"] = "Location"
    ws["B3"] = "Code"
    for r in range(4, 4 + stale_rows):
        ws.cell(row=r, column=1, value=f"OLD{r}")
        ws.cell(row=r, column=5, value=r)
    ws.merge_cells("G6:H7")
    wb.create_sheet("Other")
    wb.save(path)


def prepare_site_dir(root: Path, scenario: Scenario, include_manual_master: bool = True,
                     with_fefo_template: bool = True) -> Path:
    input_dir = root / "Data_Input"
    input_dir.mkdir(parents=True, exist_ok=True)
    write_shelf_life(input_dir / "Shelflife.xlsx", include_manual_master)
    write_capacity(input_dir / "Capacity.xlsx", scenario.locations)
    (input_dir / "SAP_MA_BO_SUNG.txt").write_text(
        "# ma bo sung\n" + ",".join(scenario.manual_materials) + "\n", encoding="utf-8"
    )
    if with_fefo_template:
        write_fefo_template(root / "FEFO.xlsx")
    return root


def scenario_from_real_inputs(seed: int, today: datetime.date, shelf_life: Path, capacity: Path,
                              n_records: int = 3000, n_stock_rows: int = 500) -> Scenario:
    rng = random.Random(seed)
    df_sl = pd.read_excel(shelf_life)
    materials = [str(v) for v in df_sl["Iteam"].dropna().astype("int64").tolist()]
    locations = pd.read_excel(capacity, sheet_name="Capacity")["Location"].dropna().astype(str).tolist()
    locations += ["K12", "H05", "LOOSE-01", "KX1"]
    records = build_records(rng, today, locations, n_records)
    for record in records:
        record["Material"] = rng.choice(materials + [UNKNOWN_MATERIAL])
    csv_text = build_stock_count(rng, today, records, locations, n_stock_rows)
    manual = [m for m in materials[-2:]]
    return Scenario(today=today, records=records, stock_count_csv=csv_text, manual_materials=manual,
                    locations=locations)
