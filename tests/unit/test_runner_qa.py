from __future__ import annotations

import datetime
import json
from dataclasses import replace

import pandas as pd
import requests
from freezegun import freeze_time

from efefo.config import load_site
from tests.conftest import FROZEN_NOW
from tests.support.fakes import FakeHttp
from tests.support.runner_harness import FEFO_URL, run_new
from tests.support.scenario import build_scenario, mutate_records, prepare_site_dir


def setup(tmp_path, seed=1, **kwargs):
    scenario = build_scenario(seed, FROZEN_NOW.date())
    root = prepare_site_dir(tmp_path / "site", scenario, **kwargs)
    return scenario, root


def test_qa01_first_run_creates_every_output(tmp_path):
    scenario, root = setup(tmp_path)
    http = FakeHttp(scenario.stock_count_csv)
    with freeze_time(FROZEN_NOW):
        result = run_new(root, tmp_path / "s", scenario.records, http)
    for name in ["Inventory_Report_Full.xlsx", "Inventory_Reconciliation_Report.xlsx", "FEFO.xlsx",
                 "sap_response_full.json", ".sap_location_history.json", "audit/audit_state.json",
                 "Daily_Reports/Inventory_Report_Full_2026-10-19.xlsx", "logs/sap_inventory.log"]:
        assert (root / name).exists(), name
    assert not (root / ".sap_inventory.lock").exists()
    assert [url for url, _ in http.posts][0] == FEFO_URL
    assert result.summary is not None and len(result.summary) > 0


def test_qa02_daily_snapshot_not_overwritten_same_day(tmp_path):
    scenario, root = setup(tmp_path)
    daily = root / "Daily_Reports" / "Inventory_Report_Full_2026-10-19.xlsx"
    with freeze_time(FROZEN_NOW):
        run_new(root, tmp_path / "s", scenario.records, FakeHttp(scenario.stock_count_csv))
    first = pd.read_excel(daily)
    with freeze_time(FROZEN_NOW + datetime.timedelta(hours=2)):
        run_new(root, tmp_path / "s", scenario.records[:50], FakeHttp(scenario.stock_count_csv))
    pd.testing.assert_frame_equal(first, pd.read_excel(daily))
    assert len(pd.read_excel(root / "Inventory_Report_Full.xlsx")) < len(first)


def test_qa03_daily_snapshot_identical_to_latest(tmp_path):
    scenario, root = setup(tmp_path)
    with freeze_time(FROZEN_NOW):
        run_new(root, tmp_path / "s", scenario.records, FakeHttp(scenario.stock_count_csv))
    pd.testing.assert_frame_equal(pd.read_excel(root / "Inventory_Report_Full.xlsx"),
                                  pd.read_excel(root / "Daily_Reports" / "Inventory_Report_Full_2026-10-19.xlsx"))


def test_qa04_no_records_stops_without_outputs(tmp_path):
    scenario, root = setup(tmp_path)
    with freeze_time(FROZEN_NOW):
        assert run_new(root, tmp_path / "s", [], FakeHttp(scenario.stock_count_csv)) is None
    assert not (root / "Inventory_Report_Full.xlsx").exists()
    assert not (root / ".sap_inventory.lock").exists()


def test_qa05_active_lock_skips_run(tmp_path):
    scenario, root = setup(tmp_path)
    (root / ".sap_inventory.lock").write_text("busy")
    with freeze_time(FROZEN_NOW):
        assert run_new(root, tmp_path / "s", scenario.records, FakeHttp(scenario.stock_count_csv)) is None
    assert (root / ".sap_inventory.lock").exists()


def test_qa06_missing_template_still_publishes(tmp_path):
    scenario, root = setup(tmp_path, with_fefo_template=False)
    http = FakeHttp(scenario.stock_count_csv)
    with freeze_time(FROZEN_NOW):
        run_new(root, tmp_path / "s", scenario.records, http)
    assert not (root / "FEFO.xlsx").exists()
    assert http.posts


def test_qa07_publish_switches_off_means_no_network(tmp_path):
    scenario, root = setup(tmp_path)
    site = load_site("hm")
    site = replace(site, publish=replace(site.publish, fefo_to_google_sheet=False,
                                         empty_locations_to_google_sheet=False, audit_to_google_sheet=False))
    http = FakeHttp(scenario.stock_count_csv)
    with freeze_time(FROZEN_NOW):
        run_new(root, tmp_path / "s", scenario.records, http, site=site)
    assert http.posts == [] and (root / "Inventory_Report_Full.xlsx").exists()


def test_qa08_network_down_keeps_reports_and_queues_audit(tmp_path):
    scenario, root = setup(tmp_path, seed=101)
    site = load_site("hm")
    site = replace(site, audit=replace(site.audit, confirm_runs=1))
    second = mutate_records(scenario.records, 5, drop_ratio=0.0)
    with freeze_time(FROZEN_NOW):
        run_new(root, tmp_path / "s", scenario.records, FakeHttp(scenario.stock_count_csv, fail_posts=True), site=site)
    with freeze_time(FROZEN_NOW + datetime.timedelta(hours=1)):
        run_new(root, tmp_path / "s", second, FakeHttp(scenario.stock_count_csv, fail_posts=True), site=site)
    state = json.loads((root / "audit" / "audit_state.json").read_text(encoding="utf-8"))
    assert state["unsent"] and all(item["attempts"] == 1 for item in state["unsent"])
    assert (root / "Inventory_Report_Full.xlsx").exists()


def test_qa09_stock_count_unreachable(tmp_path):
    scenario, root = setup(tmp_path)
    http = FakeHttp(get_error=requests.ConnectionError("dns"))
    with freeze_time(FROZEN_NOW):
        result = run_new(root, tmp_path / "s", scenario.records, http)
    assert result.stock_count is None and result.blocked == set()
    assert not (root / "Inventory_Reconciliation_Report.xlsx").exists()
    assert (root / "FEFO.xlsx").exists()


def test_qa10_fefo_workbook_layout(tmp_path):
    scenario, root = setup(tmp_path)
    with freeze_time(FROZEN_NOW):
        result = run_new(root, tmp_path / "s", scenario.records, FakeHttp(scenario.stock_count_csv))
    from openpyxl import load_workbook

    ws = load_workbook(root / "FEFO.xlsx")["Update Actual Count"]
    assert ws["A1"].value == "FEFO" and ws["A3"].value == "Location"
    rows = [r for r in ws.iter_rows(min_row=4, max_col=6, values_only=True) if any(v is not None for v in r)]
    assert len(rows) == len(result.summary)
    assert not any(str(r[0]).startswith("OLD") for r in rows)


def test_qa11_fefo_payload_contract(tmp_path):
    scenario, root = setup(tmp_path)
    http = FakeHttp(scenario.stock_count_csv)
    with freeze_time(FROZEN_NOW):
        run_new(root, tmp_path / "s", scenario.records, http)
    payload = http.posts[0][1]
    assert payload["header"] == ["Location_1", "Code", "SKU_2", "PRD date_3", "Quantity",
                                 "Remaining shelf-life (%)_4", "", "", "", "2026/10/19 08:00:00"]
    assert all(len(r) == 10 and r[5].endswith("%") for r in payload["rows"] if r[5])
    pct = [float(r[5][:-1]) for r in payload["rows"] if r[5]]
    assert pct == sorted(pct)


def test_qa12_manual_rows_never_outrank_fresher_sap_rows(tmp_path):
    scenario, root = setup(tmp_path, seed=8)
    with freeze_time(FROZEN_NOW):
        result = run_new(root, tmp_path / "s", scenario.records, FakeHttp(scenario.stock_count_csv))
    summary = result.summary
    assert summary["phantram_shelf_life"].is_monotonic_increasing
    for _, group in summary.groupby("phantram_shelf_life"):
        manual = group["Normalized_Location"].str.endswith("*").tolist()
        assert manual == sorted(manual)
